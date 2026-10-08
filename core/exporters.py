"""Excel va PDF eksport yordamchilari."""
import io
from datetime import datetime
from xml.sax.saxutils import escape

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FONTS = {}


def get_fonts():
    """Unicode (o‘, g‘, kirill) belgilarni qo'llovchi shriftlarni ro'yxatdan o'tkazadi."""
    if _FONTS:
        return _FONTS["regular"], _FONTS["bold"]
    font_dir = settings.BASE_DIR / "static" / "fonts"
    regular_path, bold_path = font_dir / "DejaVuSans.ttf", font_dir / "DejaVuSans-Bold.ttf"
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", str(regular_path)))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(bold_path)))
        _FONTS.update(regular="DejaVu", bold="DejaVu-Bold")
    except Exception:
        _FONTS.update(regular="Helvetica", bold="Helvetica-Bold")
    return _FONTS["regular"], _FONTS["bold"]


def pdf_text(value):
    """Matnni reportlab Paragraph uchun xavfsiz ko'rinishga keltiradi."""
    if value is None:
        return ""
    return escape(str(value)).replace("\n", "<br/>")


def _safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value).replace(tzinfo=None)
    return value


def _stamp():
    return timezone.localtime().strftime("%Y%m%d_%H%M")


def xlsx_response(title, headers, rows, filename_prefix="hisobot", subtitle=""):
    wb = Workbook()
    ws = wb.active
    ws.title = (title[:28] or "Hisobot").replace("/", "-").replace("\\", "-").replace("*", "").replace("?", "").replace(":", "").replace("[", "").replace("]", "")
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle or f"Shakllantirilgan: {timezone.localtime():%d.%m.%Y %H:%M}"])
    ws.append([])
    ws.append(list(headers))
    header_row = 4
    fill = PatternFill("solid", fgColor="1F4E79")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=header_row, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_safe_cell(v) for v in row])
    for idx, header in enumerate(headers, start=1):
        width = len(str(header))
        for row in rows:
            if idx - 1 < len(row):
                width = max(width, min(len(str(row[idx - 1])), 60))
        ws.column_dimensions[get_column_letter(idx)].width = width + 3
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type=XLSX_MIME)
    response["Content-Disposition"] = f'attachment; filename="{filename_prefix}_{_stamp()}.xlsx"'
    return response


def pdf_table_bytes(title, headers, rows, subtitle=""):
    regular, bold = get_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Title"], fontName=bold, fontSize=15, alignment=0)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontName=regular, fontSize=9, textColor=colors.grey)
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontName=regular, fontSize=8, leading=10)
    head = ParagraphStyle("head", parent=cell, fontName=bold, textColor=colors.white)
    story = [
        Paragraph(pdf_text(title), h),
        Paragraph(pdf_text(subtitle or f"Samarqand iqtisodiyot va servis instituti — Karyera markazi. Shakllantirilgan: {timezone.localtime():%d.%m.%Y %H:%M}"), sub),
        Spacer(1, 6 * mm),
    ]
    data = [[Paragraph(pdf_text(x), head) for x in headers]]
    for row in rows:
        data.append([Paragraph(pdf_text("" if v is None else v), cell) for v in row])
    if len(data) == 1:
        data.append([Paragraph("Ma'lumot topilmadi", cell)] + [""] * (len(headers) - 1))
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B0B7C3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    doc.build(story)
    return buf.getvalue()


def pdf_response(title, headers, rows, filename_prefix="hisobot", subtitle=""):
    content = pdf_table_bytes(title, headers, rows, subtitle)
    response = HttpResponse(content, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename_prefix}_{_stamp()}.pdf"'
    return response
