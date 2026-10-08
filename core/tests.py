import io
from datetime import timedelta

from django.core import mail
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import Direction, Faculty, Region, Role, User
from core.models import AuditLog, Notification
from employers.models import Employer
from events.models import Event, EventRegistration
from reports.registry import REPORTS
from students.models import EmploymentStatus, StudentProfile
from vacancies.models import Application, Vacancy

PWD = "Str0ng!Pass2026"


def make_user(username, role, faculty=None, **kw):
    u = User.objects.create_user(username=username, password=PWD, role=role, faculty=faculty,
                                 first_name=kw.pop("first_name", username.title()), last_name=kw.pop("last_name", "Test"),
                                 email=f"{username}@example.uz", **kw)
    return u


class BaseCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.f1 = Faculty.objects.create(name="Iqtisodiyot")
        cls.f2 = Faculty.objects.create(name="Servis")
        cls.d1 = Direction.objects.create(faculty=cls.f1, name="Buxgalteriya")
        cls.d2 = Direction.objects.create(faculty=cls.f2, name="Turizm")
        cls.admin = make_user("admin1", Role.ADMIN)
        cls.staff = make_user("staff1", Role.STAFF)
        cls.dean = make_user("dean1", Role.DEAN, faculty=cls.f1)
        cls.mgmt = make_user("mgmt1", Role.MANAGEMENT)
        cls.student = make_user("student1", Role.STUDENT)
        cls.grad = make_user("grad1", Role.GRADUATE)
        cls.grad2 = make_user("grad2", Role.GRADUATE)
        cls.p_student = StudentProfile.objects.create(user=cls.student, faculty=cls.f1, direction=cls.d1, course=2)
        cls.p_grad = StudentProfile.objects.create(user=cls.grad, faculty=cls.f1, direction=cls.d1, graduation_year=2025,
                                                   employment_status=EmploymentStatus.EMPLOYED, employer_name="Bank",
                                                   position="Buxgalter", employed_date=timezone.localdate())
        cls.p_grad2 = StudentProfile.objects.create(user=cls.grad2, faculty=cls.f2, direction=cls.d2, graduation_year=2025,
                                                    employment_status=EmploymentStatus.LOOKING)
        cls.emp_user = make_user("emp1", Role.EMPLOYER)
        cls.employer = Employer.objects.create(user=cls.emp_user, name="Bank", stir="123456789", industry="Moliya",
                                               region=Region.SAMARKAND, address="Samarqand", phone="+998662345678",
                                               email="b@b.uz", contact_name="X", contact_phone="+998901234567",
                                               status=Employer.Status.APPROVED)
        cls.emp2_user = make_user("emp2", Role.EMPLOYER)
        cls.employer2 = Employer.objects.create(user=cls.emp2_user, name="Other", stir="987654321", industry="IT",
                                                region=Region.SAMARKAND, address="Samarqand", phone="+998662345679",
                                                email="o@o.uz", contact_name="Y", contact_phone="+998901234568",
                                                status=Employer.Status.APPROVED)
        cls.vacancy = Vacancy.objects.create(employer=cls.employer, title="Buxgalter", region=Region.SAMARKAND,
                                             requirements="req", duties="duty", deadline=timezone.localdate() + timedelta(days=10),
                                             status=Vacancy.Status.PUBLISHED, published_at=timezone.now(), salary_min=3_000_000)
        cls.event = Event.objects.create(title="Yarmarka", event_type="job_fair", description="d", location="Zal",
                                         start_at=timezone.now() + timedelta(days=5), capacity=2)

    def login(self, user):
        self.client.force_login(user)


class PublicPagesTests(BaseCase):
    def test_public_pages_ok(self):
        for name, args in [("core:home", []), ("vacancies:list", []), ("vacancies:detail", [self.vacancy.pk]),
                           ("employers:public_list", []), ("employers:public_detail", [self.employer.pk]),
                           ("events:list", []), ("events:detail", [self.event.pk]), ("core:news_list", []),
                           ("accounts:login", []), ("accounts:register", []), ("accounts:register_employer", []),
                           ("accounts:password_reset", []), ("core:health", [])]:
            r = self.client.get(reverse(name, args=args))
            self.assertEqual(r.status_code, 200, name)

    def test_protected_pages_redirect_anonymous(self):
        for name in ["students:profile", "vacancies:my_applications", "employers:dashboard", "reports:dashboard",
                     "students:monitor_list", "accounts:user_list", "core:audit_list", "vacancies:manage_list"]:
            r = self.client.get(reverse(name))
            self.assertEqual(r.status_code, 302, name)
            self.assertIn("/hisob/kirish/", r["Location"])

    def test_vacancy_filters(self):
        r = self.client.get(reverse("vacancies:list"), {"q": "Buxg", "salary": "2000000", "direction": self.d1.pk,
                                                       "region": Region.SAMARKAND, "type": "full", "sort": "salary"})
        self.assertContains(r, "Buxgalter")
        r = self.client.get(reverse("vacancies:list"), {"q": "yo'q-bunaqa"})
        self.assertEqual(r.context["page"].paginator.count, 0)
        r = self.client.get(reverse("vacancies:list"), {"salary": "99000000"})
        self.assertEqual(r.context["page"].paginator.count, 0)

    def test_unpublished_vacancy_hidden(self):
        v = Vacancy.objects.create(employer=self.employer, title="Qoralama", region=Region.SAMARKAND, requirements="r", duties="d",
                                   deadline=timezone.localdate() + timedelta(days=3), status=Vacancy.Status.DRAFT)
        self.assertEqual(self.client.get(reverse("vacancies:detail", args=[v.pk])).status_code, 404)
        self.login(self.emp_user)
        self.assertEqual(self.client.get(reverse("vacancies:detail", args=[v.pk])).status_code, 200)

    def test_health(self):
        self.assertEqual(self.client.get(reverse("core:health")).json()["status"], "ok")


class AuthTests(BaseCase):
    def test_login_logout_audit(self):
        r = self.client.post(reverse("accounts:login"), {"username": "student1", "password": PWD})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(AuditLog.objects.filter(action="login", username="student1").exists())
        self.client.post(reverse("accounts:logout"))
        self.assertTrue(AuditLog.objects.filter(action="logout", username="student1").exists())

    @override_settings(LOGIN_MAX_ATTEMPTS=3)
    def test_lockout_after_failed_attempts(self):
        cache.clear()
        for _ in range(3):
            self.client.post(reverse("accounts:login"), {"username": "student1", "password": "bad"})
        self.assertEqual(AuditLog.objects.filter(action="login_failed").count(), 3)
        r = self.client.post(reverse("accounts:login"), {"username": "student1", "password": PWD})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Juda ko&#x27;p")
        cache.clear()

    def test_role_redirects(self):
        for user, target in [(self.student, "/kabinet/profil/"), (self.emp_user, "/ish-beruvchi/"), (self.staff, "/dashboard/"),
                             (self.mgmt, "/dashboard/"), (self.dean, "/dashboard/"), (self.admin, "/dashboard/")]:
            self.login(user)
            r = self.client.get(reverse("core:redirect_by_role"))
            self.assertEqual(r["Location"], target)

    def test_candidate_registration(self):
        data = {"role": "student", "last_name": "Aliyev", "first_name": "Vali", "patronymic": "Akramovich", "faculty": self.f1.pk,
                "direction": self.d1.pk, "course": 2, "email": "vali@example.uz", "phone": "+998 90 123-45-67",
                "username": "vali", "password1": PWD, "password2": PWD}
        r = self.client.post(reverse("accounts:register"), data)
        self.assertEqual(r.status_code, 302, getattr(r, "context", None) and r.context["form"].errors)
        u = User.objects.get(username="vali")
        self.assertEqual(u.phone, "+998901234567")
        self.assertEqual(u.profile.course, 2)

    def test_registration_validation(self):
        data = {"role": "graduate", "last_name": "A", "first_name": "B", "faculty": self.f1.pk, "direction": self.d2.pk,
                "email": "student1@example.uz", "phone": "123", "username": "x1", "password1": PWD, "password2": PWD}
        r = self.client.post(reverse("accounts:register"), data)
        self.assertEqual(r.status_code, 200)
        errors = r.context["form"].errors
        for field in ("direction", "graduation_year", "email", "phone"):
            self.assertIn(field, errors)
        self.assertFalse(User.objects.filter(username="x1").exists())

    def test_employer_registration(self):
        data = {"name": "Yangi MCHJ", "stir": "555666777", "industry": "Savdo", "region": Region.SAMARKAND, "address": "Manzil",
                "org_phone": "+998662000000", "last_name": "Rahimov", "first_name": "Karim", "email": "new@emp.uz",
                "phone": "+998901112233", "username": "newemp", "password1": PWD, "password2": PWD}
        self.assertEqual(self.client.post(reverse("accounts:register_employer"), data).status_code, 302)
        e = Employer.objects.get(stir="555666777")
        self.assertEqual(e.status, Employer.Status.PENDING)
        self.assertTrue(Notification.objects.filter(user=self.staff, title="Yangi ish beruvchi").exists())
        # kirgan foydalanuvchi qayta ro'yxatdan o'tmaydi
        self.assertEqual(self.client.get(reverse("accounts:register_employer")).status_code, 302)
        self.client.logout()
        dup = self.client.post(reverse("accounts:register_employer"), dict(data, username="other", email="o2@emp.uz"))
        self.assertEqual(dup.status_code, 200)
        self.assertIn("stir", dup.context["form"].errors)

    def test_admin_role_sets_staff_flags(self):
        self.assertTrue(self.admin.is_superuser and self.admin.is_staff)
        self.assertFalse(self.staff.is_staff or self.staff.is_superuser)

    def test_password_reset_sends_email(self):
        r = self.client.post(reverse("accounts:password_reset"), {"email": "student1@example.uz"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("/hisob/parol-tiklash/", mail.outbox[0].body)


class PermissionTests(BaseCase):
    def test_student_cannot_access_staff_areas(self):
        self.login(self.student)
        for name in ["students:monitor_list", "reports:dashboard", "reports:index", "vacancies:manage_list", "employers:manage_list",
                     "accounts:user_list", "core:audit_list", "employers:dashboard", "vacancies:employer_list",
                     "core:faculty_list", "events:manage_list"]:
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)

    def test_employer_cannot_access_candidate_or_staff(self):
        self.login(self.emp_user)
        for name in ["students:profile", "vacancies:my_applications", "reports:dashboard", "students:monitor_list", "vacancies:manage_list"]:
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)

    def test_only_admin_manages_users_and_audit(self):
        for u in (self.staff, self.mgmt, self.dean):
            self.login(u)
            self.assertEqual(self.client.get(reverse("accounts:user_list")).status_code, 403)
            self.assertEqual(self.client.get(reverse("core:audit_list")).status_code, 403)
        self.login(self.admin)
        self.assertEqual(self.client.get(reverse("accounts:user_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("core:audit_list")).status_code, 200)

    def test_management_is_read_only(self):
        self.login(self.mgmt)
        self.assertEqual(self.client.get(reverse("students:monitor_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("students:monitor_edit", args=[self.p_grad.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("vacancies:manage_list")).status_code, 403)

    def test_dean_sees_only_own_faculty(self):
        self.login(self.dean)
        r = self.client.get(reverse("students:monitor_list"))
        names = [p.user.username for p in r.context["page"]]
        self.assertIn("grad1", names)
        self.assertNotIn("grad2", names)
        self.assertEqual(self.client.get(reverse("students:monitor_detail", args=[self.p_grad2.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("students:cv_other", args=[self.p_grad2.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("students:cv_other", args=[self.p_grad.pk])).status_code, 200)

    def test_employer_sees_cv_only_of_applicants(self):
        self.login(self.emp_user)
        self.assertEqual(self.client.get(reverse("students:cv_other", args=[self.p_grad.pk])).status_code, 403)
        Application.objects.create(vacancy=self.vacancy, applicant=self.grad)
        self.assertEqual(self.client.get(reverse("students:cv_other", args=[self.p_grad.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("students:cv_pdf_other", args=[self.p_grad.pk])).status_code, 200)
        self.login(self.emp2_user)
        self.assertEqual(self.client.get(reverse("students:cv_other", args=[self.p_grad.pk])).status_code, 403)

    def test_employer_cannot_touch_other_employers_data(self):
        app = Application.objects.create(vacancy=self.vacancy, applicant=self.grad)
        self.login(self.emp2_user)
        self.assertEqual(self.client.get(reverse("vacancies:employer_application", args=[app.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("vacancies:employer_edit", args=[self.vacancy.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse("vacancies:employer_close", args=[self.vacancy.pk])).status_code, 404)

    def test_admin_cannot_demote_self(self):
        self.login(self.admin)
        r = self.client.post(reverse("accounts:user_edit", args=[self.admin.pk]), {
            "last_name": "T", "first_name": "A", "email": "a@a.uz", "role": "student", "is_active": ""})
        self.admin.refresh_from_db()
        self.assertEqual(self.admin.role, Role.ADMIN)
        self.assertTrue(self.admin.is_active)

    def test_role_change_is_audited(self):
        self.login(self.admin)
        self.client.post(reverse("accounts:user_edit", args=[self.staff.pk]), {
            "last_name": "T", "first_name": "S", "email": "s@s.uz", "role": "management", "is_active": "on"})
        self.staff.refresh_from_db()
        self.assertEqual(self.staff.role, Role.MANAGEMENT)
        self.assertTrue(AuditLog.objects.filter(action="role_change", object_id=str(self.staff.pk)).exists())


class VacancyWorkflowTests(BaseCase):
    def vacancy_data(self, **kw):
        d = {"title": "Dasturchi", "region": Region.SAMARKAND, "employment_type": "full", "salary_min": 5000000, "salary_max": 8000000,
             "schedule": "9-18", "deadline": (timezone.localdate() + timedelta(days=20)).isoformat(), "directions": [self.d1.pk],
             "requirements": "Python", "duties": "Kod yozish", "action": "submit"}
        d.update(kw)
        return d

    def test_full_lifecycle(self):
        # ish beruvchi vakansiya yaratadi va moderatsiyaga yuboradi
        self.login(self.emp_user)
        r = self.client.post(reverse("vacancies:employer_create"), self.vacancy_data())
        self.assertEqual(r.status_code, 302)
        v = Vacancy.objects.get(title="Dasturchi")
        self.assertEqual(v.status, Vacancy.Status.PENDING)
        self.assertTrue(AuditLog.objects.filter(action="vacancy_create").exists())
        self.assertTrue(Notification.objects.filter(user=self.staff, title__icontains="moderatsiya").exists())
        self.assertEqual(self.client.get(reverse("vacancies:detail", args=[v.pk])).status_code, 200)  # egasi ko'ra oladi
        # moderator rad etadi (sababsiz — xato)
        self.login(self.staff)
        r = self.client.post(reverse("vacancies:manage_detail", args=[v.pk]), {"decision": "reject", "comment": ""})
        self.assertEqual(r.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.status, Vacancy.Status.PENDING)
        # sabab bilan rad etadi
        self.client.post(reverse("vacancies:manage_detail", args=[v.pk]), {"decision": "reject", "comment": "To'ldiring"})
        v.refresh_from_db()
        self.assertEqual(v.status, Vacancy.Status.REJECTED)
        # ish beruvchi tahrirlab qayta yuboradi
        self.login(self.emp_user)
        self.client.post(reverse("vacancies:employer_edit", args=[v.pk]), self.vacancy_data(title="Dasturchi (Django)"))
        v.refresh_from_db()
        self.assertEqual(v.status, Vacancy.Status.PENDING)
        # moderator tasdiqlaydi
        self.login(self.staff)
        self.client.post(reverse("vacancies:manage_detail", args=[v.pk]), {"decision": "publish", "comment": ""})
        v.refresh_from_db()
        self.assertEqual(v.status, Vacancy.Status.PUBLISHED)
        self.assertIsNotNone(v.published_at)
        self.assertTrue(Notification.objects.filter(user=self.student, title="Yangi vakansiya").exists())
        # talaba ariza topshiradi
        self.login(self.student)
        r = self.client.post(reverse("vacancies:apply", args=[v.pk]), {"cover_letter": "Salom"})
        self.assertEqual(r.status_code, 302)
        app = Application.objects.get(vacancy=v, applicant=self.student)
        self.assertTrue(Notification.objects.filter(user=self.emp_user, title="Yangi ariza").exists())
        # ikkinchi marta topshirib bo'lmaydi
        self.client.post(reverse("vacancies:apply", args=[v.pk]), {"cover_letter": "Yana"})
        self.assertEqual(Application.objects.filter(vacancy=v, applicant=self.student).count(), 1)
        # ish beruvchi suhbatga taklif qiladi
        self.login(self.emp_user)
        r = self.client.post(reverse("vacancies:employer_application", args=[app.pk]), {"status": "interview", "interview_place": "Ofis"})
        self.assertEqual(r.status_code, 200)  # suhbat vaqti majburiy
        when = (timezone.localtime() + timedelta(days=3)).strftime("%Y-%m-%dT%H:%M")
        r = self.client.post(reverse("vacancies:employer_application", args=[app.pk]),
                             {"status": "interview", "interview_at": when, "interview_place": "Ofis", "employer_comment": "Kelishilgan"})
        self.assertEqual(r.status_code, 302)
        app.refresh_from_db()
        self.assertEqual(app.status, Application.Status.INTERVIEW)
        self.assertTrue(Notification.objects.filter(user=self.student, title="Ariza holati o'zgardi").exists())
        # talaba arizani bekor qiladi va qayta topshiradi
        self.login(self.student)
        self.client.post(reverse("vacancies:withdraw", args=[app.pk]))
        app.refresh_from_db()
        self.assertEqual(app.status, Application.Status.WITHDRAWN)
        self.client.post(reverse("vacancies:apply", args=[v.pk]), {"cover_letter": "Qayta"})
        app.refresh_from_db()
        self.assertEqual(app.status, Application.Status.SENT)

    def test_employer_must_be_approved(self):
        self.employer.status = Employer.Status.PENDING
        self.employer.save()
        self.login(self.emp_user)
        r = self.client.post(reverse("vacancies:employer_create"), self.vacancy_data())
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Vacancy.objects.filter(title="Dasturchi").exists())
        self.assertEqual(self.client.get(reverse("vacancies:detail", args=[self.vacancy.pk])).status_code, 200)  # egasi
        self.client.logout()
        self.assertEqual(self.client.get(reverse("vacancies:detail", args=[self.vacancy.pk])).status_code, 404)

    def test_validation_rules(self):
        self.login(self.emp_user)
        r = self.client.post(reverse("vacancies:employer_create"),
                             self.vacancy_data(deadline=(timezone.localdate() - timedelta(days=1)).isoformat(), salary_min=9000000, salary_max=1000000))
        self.assertEqual(r.status_code, 200)
        self.assertIn("deadline", r.context["form"].errors)
        self.assertIn("salary_max", r.context["form"].errors)

    def test_expired_vacancy_cannot_receive_applications(self):
        self.vacancy.deadline = timezone.localdate() - timedelta(days=1)
        self.vacancy.save()
        self.login(self.student)
        self.client.post(reverse("vacancies:apply", args=[self.vacancy.pk]), {})
        self.assertFalse(Application.objects.exists())
        self.assertEqual(self.client.get(reverse("vacancies:list")).context["page"].paginator.count, 0)

    def test_non_candidates_cannot_apply(self):
        self.login(self.staff)
        self.assertEqual(self.client.post(reverse("vacancies:apply", args=[self.vacancy.pk])).status_code, 403)

    def test_delete_rules(self):
        self.login(self.emp_user)
        self.client.post(reverse("vacancies:employer_delete", args=[self.vacancy.pk]))
        self.assertTrue(Vacancy.objects.filter(pk=self.vacancy.pk).exists())  # e'lon qilingan — o'chirilmaydi
        v = Vacancy.objects.create(employer=self.employer, title="Draft", region=Region.SAMARKAND, requirements="r", duties="d",
                                   deadline=timezone.localdate() + timedelta(days=5))
        self.client.post(reverse("vacancies:employer_delete", args=[v.pk]))
        self.assertFalse(Vacancy.objects.filter(pk=v.pk).exists())
        self.assertTrue(AuditLog.objects.filter(action="vacancy_delete").exists())

    def test_employer_status_notifies(self):
        self.login(self.staff)
        r = self.client.post(reverse("employers:manage_detail", args=[self.employer.pk]),
                             {"status": "blocked", "status_comment": "", "cooperation": "none"})
        self.assertEqual(r.status_code, 200)  # sabab majburiy
        self.client.post(reverse("employers:manage_detail", args=[self.employer.pk]),
                         {"status": "blocked", "status_comment": "Qoidabuzarlik", "cooperation": "none"})
        self.employer.refresh_from_db()
        self.assertEqual(self.employer.status, Employer.Status.BLOCKED)
        self.assertEqual(self.client.get(reverse("vacancies:list")).context["page"].paginator.count, 0)
        self.assertTrue(Notification.objects.filter(user=self.emp_user, title="Tashkilot holati o'zgardi").exists())


class StudentProfileTests(BaseCase):
    def test_profile_edit_and_employment_rules(self):
        self.login(self.grad2)
        data = {"faculty": self.f2.pk, "direction": self.d2.pk, "graduation_year": 2025, "employment_status": "employed",
                "employer_name": "", "position": ""}
        r = self.client.post(reverse("students:profile_edit"), data)
        self.assertEqual(r.status_code, 200)
        data.update(employer_name="IT Park", position="Dasturchi")
        self.assertEqual(self.client.post(reverse("students:profile_edit"), data).status_code, 302)
        self.p_grad2.refresh_from_db()
        self.assertEqual(self.p_grad2.employment_status, "employed")
        self.assertEqual(self.p_grad2.employed_date, timezone.localdate())
        data.update(employment_status="looking")
        self.client.post(reverse("students:profile_edit"), data)
        self.p_grad2.refresh_from_db()
        self.assertEqual(self.p_grad2.employer_name, "")
        self.assertIsNone(self.p_grad2.employed_date)

    def test_direction_must_match_faculty(self):
        self.login(self.student)
        r = self.client.post(reverse("students:profile_edit"), {"faculty": self.f1.pk, "direction": self.d2.pk, "course": 2, "employment_status": "unknown"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("direction", r.context["form"].errors)

    def test_sections_crud_and_ownership(self):
        self.login(self.student)
        self.assertEqual(self.client.post(reverse("students:section_add", args=["skills"]), {"name": "Excel", "level": 3}).status_code, 302)
        skill = self.p_student.skills.get()
        self.assertEqual(self.client.post(reverse("students:section_edit", args=["skills", skill.pk]), {"name": "Excel Pro", "level": 4}).status_code, 302)
        skill.refresh_from_db()
        self.assertEqual(skill.name, "Excel Pro")
        self.client.post(reverse("students:section_add", args=["experiences"]), {"company": "A", "position": "B", "start_date": "2024-01-01", "end_date": "2023-01-01"})
        self.assertFalse(self.p_student.experiences.exists())  # tugash sanasi xato
        self.client.post(reverse("students:section_add", args=["experiences"]), {"company": "A", "position": "B", "start_date": "2024-01-01"})
        self.assertTrue(self.p_student.experiences.exists())
        # boshqa foydalanuvchi elementiga tegib bo'lmaydi
        self.login(self.grad)
        self.assertEqual(self.client.post(reverse("students:section_delete", args=["skills", skill.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("students:section_add", args=["hacker"])).status_code, 404)
        self.login(self.student)
        self.client.post(reverse("students:section_delete", args=["skills", skill.pk]))
        self.assertFalse(self.p_student.skills.exists())

    def test_cv_html_and_pdf(self):
        self.p_student.about = "O'zim haqida <b>teg</b> & maxsus belgilar: o‘ g‘ ʻ"
        self.p_student.save()
        self.p_student.skills.create(name="Python", level=3)
        self.p_student.experiences.create(company="Kompaniya", position="Stajyor", start_date="2024-01-01")
        self.login(self.student)
        self.assertContains(self.client.get(reverse("students:cv")), "Student1")
        r = self.client.get(reverse("students:cv_pdf"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_profile_completeness(self):
        self.assertLess(self.p_student.completeness, 100)
        self.assertGreaterEqual(self.p_student.completeness, 0)

    def test_profile_created_lazily_for_role_without_profile(self):
        u = make_user("late", Role.STUDENT)
        self.login(u)
        self.assertEqual(self.client.get(reverse("students:profile")).status_code, 200)
        self.assertTrue(StudentProfile.objects.filter(user=u).exists())


class EventTests(BaseCase):
    def test_registration_capacity_and_cancel(self):
        self.login(self.student)
        self.client.post(reverse("events:register", args=[self.event.pk]))
        self.client.post(reverse("events:register", args=[self.event.pk]))
        self.assertEqual(EventRegistration.objects.filter(event=self.event).count(), 1)
        self.login(self.grad)
        self.client.post(reverse("events:register", args=[self.event.pk]))
        self.login(self.grad2)
        self.client.post(reverse("events:register", args=[self.event.pk]))  # joy yo'q (sig'im 2)
        self.assertEqual(EventRegistration.objects.filter(event=self.event).count(), 2)
        self.login(self.student)
        self.client.post(reverse("events:cancel", args=[self.event.pk]))
        self.assertEqual(EventRegistration.objects.filter(event=self.event).count(), 1)

    def test_past_event_not_registrable(self):
        past = Event.objects.create(title="Eski", event_type="seminar", description="d", location="x", start_at=timezone.now() - timedelta(days=2))
        self.login(self.student)
        self.client.post(reverse("events:register", args=[past.pk]))
        self.assertFalse(EventRegistration.objects.filter(event=past).exists())

    def test_staff_event_crud_and_attendance(self):
        self.login(self.staff)
        start = (timezone.localtime() + timedelta(days=9)).strftime("%Y-%m-%dT%H:%M")
        r = self.client.post(reverse("events:manage_create"), {"title": "Trening", "event_type": "training", "description": "d",
                                                               "location": "Zal", "start_at": start, "is_published": "on"})
        self.assertEqual(r.status_code, 302)
        ev = Event.objects.get(title="Trening")
        self.assertEqual(ev.created_by, self.staff)
        EventRegistration.objects.create(event=ev, user=self.student)
        reg = EventRegistration.objects.get(event=ev)
        self.client.post(reverse("events:participants", args=[ev.pk]), {"attended": [reg.pk]})
        reg.refresh_from_db()
        self.assertTrue(reg.attended)
        bad = self.client.post(reverse("events:manage_create"), {"title": "X", "event_type": "other", "description": "d", "location": "L",
                                                                 "start_at": start, "end_at": (timezone.localtime() - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M")})
        self.assertEqual(bad.status_code, 200)

    def test_staff_manage_pages(self):
        self.login(self.staff)
        for name in ["events:manage_list", "core:news_manage_list", "core:faculty_list", "core:direction_list",
                     "events:manage_create", "core:news_manage_create", "core:faculty_create", "core:direction_create"]:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)

    def test_faculty_delete_protected(self):
        self.login(self.staff)
        self.client.post(reverse("core:faculty_delete", args=[self.f1.pk]))
        self.assertTrue(Faculty.objects.filter(pk=self.f1.pk).exists())
        empty = Faculty.objects.create(name="Bo'sh")
        self.client.post(reverse("core:faculty_delete", args=[empty.pk]))
        self.assertFalse(Faculty.objects.filter(pk=empty.pk).exists())


class ReportTests(BaseCase):
    def test_dashboard_for_all_analytic_roles(self):
        for u in (self.staff, self.admin, self.mgmt, self.dean):
            self.login(u)
            r = self.client.get(reverse("reports:dashboard"))
            self.assertEqual(r.status_code, 200, u.role)
            self.assertIn("monthly", r.context["d"])

    def test_dean_dashboard_numbers_scoped(self):
        self.login(self.dean)
        d = self.client.get(reverse("reports:dashboard")).context["d"]
        self.assertEqual(d["graduates_total"], 1)
        self.login(self.staff)
        d = self.client.get(reverse("reports:dashboard")).context["d"]
        self.assertEqual(d["graduates_total"], 2)
        self.assertEqual(d["hired"], 1)
        self.assertEqual(d["employment_rate"], 50.0)

    def test_all_reports_render_and_export(self):
        self.login(self.staff)
        for key in REPORTS:
            url = reverse("reports:report", args=[key])
            self.assertEqual(self.client.get(url).status_code, 200, key)
            x = self.client.get(url, {"export": "xlsx"})
            self.assertEqual(x.status_code, 200, key)
            wb = load_workbook(io.BytesIO(x.content))
            self.assertGreaterEqual(wb.active.max_row, 4)
            p = self.client.get(url, {"export": "pdf"})
            self.assertEqual(p.status_code, 200, key)
            self.assertTrue(p.content.startswith(b"%PDF"), key)
        self.assertTrue(AuditLog.objects.filter(action="export").exists())

    def test_report_filters_and_validation(self):
        self.login(self.staff)
        r = self.client.get(reverse("reports:report", args=["by-faculty"]), {"faculty": self.f1.pk, "year": 2025})
        rows = r.context["rows"]
        self.assertEqual(rows[0][0], "Iqtisodiyot")
        self.assertEqual(rows[0][1], 1)
        r = self.client.get(reverse("reports:report", args=["dynamics"]), {"date_from": "2026-05-01", "date_to": "2026-01-01"})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.context["form"].errors)
        for period in ("month", "quarter", "year"):
            self.assertEqual(self.client.get(reverse("reports:report", args=["dynamics"]), {"period": period}).status_code, 200)

    def test_report_role_restrictions(self):
        self.login(self.dean)
        self.assertEqual(self.client.get(reverse("reports:report", args=["employers"])).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:report", args=["by-faculty"])).status_code, 200)
        self.assertEqual(self.client.get(reverse("reports:report", args=["nope"])).status_code, 404)
        rows = self.client.get(reverse("reports:report", args=["by-faculty"]), {"faculty": self.f2.pk}).context["rows"]
        self.assertEqual(rows[0][0], "Iqtisodiyot")  # dekan boshqa fakultetni tanlay olmaydi

    def test_students_export(self):
        self.login(self.dean)
        r = self.client.get(reverse("students:monitor_list"), {"export": "xlsx"})
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.active.max_row, 4 + 2)  # faqat o'z fakulteti: talaba + bitiruvchi
        self.assertEqual(self.client.get(reverse("students:monitor_list"), {"export": "pdf"}).status_code, 200)

    def test_excel_formula_injection_escaped(self):
        self.p_student.user.last_name = "=HYPERLINK(\"http://evil\")"
        self.p_student.user.save()
        self.login(self.staff)
        r = self.client.get(reverse("students:monitor_list"), {"export": "xlsx"})
        ws = load_workbook(io.BytesIO(r.content)).active
        values = [c.value for row in ws.iter_rows(min_row=5) for c in row if isinstance(c.value, str)]
        self.assertTrue(any(v.startswith("'=") for v in values))
        self.assertFalse(any(v.startswith("=") for v in values))

    def test_monitor_filters_and_edit(self):
        self.login(self.staff)
        r = self.client.get(reverse("students:monitor_list"), {"status": "looking", "role": "graduate", "year": 2025, "q": "grad2"})
        self.assertEqual([p.user.username for p in r.context["page"]], ["grad2"])
        r = self.client.post(reverse("students:monitor_edit", args=[self.p_grad2.pk]), {
            "faculty": self.f2.pk, "direction": self.d2.pk, "graduation_year": 2025, "employment_status": "employed",
            "employer_name": "Firma", "position": "Menejer", "employed_date": "2026-02-01"})
        self.assertEqual(r.status_code, 302)
        self.p_grad2.refresh_from_db()
        self.assertEqual(self.p_grad2.employment_status, "employed")
        self.assertTrue(AuditLog.objects.filter(action="data_change", object_id=str(self.p_grad2.pk)).exists())


class ApiTests(BaseCase):
    def setUp(self):
        cache.clear()
        self.api = APIClient()

    def test_public_endpoints(self):
        for path in ["/api/v1/vacancies/", f"/api/v1/vacancies/{self.vacancy.pk}/", "/api/v1/events/", "/api/v1/news/",
                     "/api/v1/faculties/", f"/api/v1/directions/?faculty={self.f1.pk}"]:
            self.assertEqual(self.api.get(path).status_code, 200, path)
        data = self.api.get("/api/v1/vacancies/", {"search": "Buxg"}).json()
        self.assertEqual(data["count"], 1)

    def test_auth_required_endpoints(self):
        for path in ["/api/v1/me/", "/api/v1/applications/", "/api/v1/stats/"]:
            self.assertIn(self.api.get(path).status_code, (401, 403), path)

    def test_token_login_and_apply(self):
        r = self.api.post("/api/v1/auth/token/", {"username": "student1", "password": PWD}, format="json")
        self.assertEqual(r.status_code, 200)
        self.api.credentials(HTTP_AUTHORIZATION=f"Token {r.json()['token']}")
        self.assertEqual(self.api.get("/api/v1/me/").json()["user"]["role"], "student")
        r = self.api.post("/api/v1/applications/", {"vacancy": self.vacancy.pk, "cover_letter": "API"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        again = self.api.post("/api/v1/applications/", {"vacancy": self.vacancy.pk}, format="json")
        self.assertEqual(again.status_code, 400)
        app_id = r.json()["id"]
        self.assertEqual(self.api.post(f"/api/v1/applications/{app_id}/withdraw/").status_code, 200)
        self.assertEqual(self.api.patch("/api/v1/me/", {"desired_position": "Analitik"}, format="json").status_code, 200)

    def test_bad_token_login(self):
        r = self.api.post("/api/v1/auth/token/", {"username": "student1", "password": "bad"}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_employer_api_sees_only_own_applications(self):
        app = Application.objects.create(vacancy=self.vacancy, applicant=self.student)
        self.api.force_authenticate(self.emp_user)
        self.assertEqual(self.api.get("/api/v1/applications/").json()["count"], 1)
        when = (timezone.now() + timedelta(days=2)).isoformat()
        r = self.api.post(f"/api/v1/applications/{app.pk}/set-status/", {"status": "interview", "interview_at": when}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(self.api.post(f"/api/v1/applications/{app.pk}/set-status/", {"status": "interview"}, format="json").status_code, 400)
        self.api.force_authenticate(self.emp2_user)
        self.assertEqual(self.api.get("/api/v1/applications/").json()["count"], 0)
        self.assertEqual(self.api.post(f"/api/v1/applications/{app.pk}/set-status/", {"status": "accepted"}, format="json").status_code, 404)
        self.api.force_authenticate(self.student)
        self.assertEqual(self.api.post(f"/api/v1/applications/{app.pk}/set-status/", {"status": "accepted"}, format="json").status_code, 403)

    def test_stats_role_based(self):
        self.api.force_authenticate(self.student)
        self.assertEqual(self.api.get("/api/v1/stats/").status_code, 403)
        self.api.force_authenticate(self.staff)
        self.assertEqual(self.api.get("/api/v1/stats/").status_code, 200)

    def test_event_register_api(self):
        self.api.force_authenticate(self.student)
        self.assertEqual(self.api.post(f"/api/v1/events/{self.event.pk}/register/").status_code, 201)
        self.assertEqual(self.api.post(f"/api/v1/events/{self.event.pk}/register/").status_code, 200)
        self.assertEqual(self.api.post(f"/api/v1/events/{self.event.pk}/cancel/").status_code, 200)
        self.api.force_authenticate(self.staff)
        self.assertEqual(self.api.post(f"/api/v1/events/{self.event.pk}/register/").status_code, 403)


class MiscTests(BaseCase):
    def test_notifications_flow(self):
        n = Notification.objects.create(user=self.student, title="T", message="M", link="/vakansiyalar/")
        other = Notification.objects.create(user=self.grad, title="X")
        self.login(self.student)
        self.assertContains(self.client.get(reverse("core:notifications")), "T")
        r = self.client.get(reverse("core:notification_open", args=[n.pk]))
        self.assertEqual(r["Location"], "/vakansiyalar/")
        n.refresh_from_db()
        self.assertTrue(n.is_read)
        self.assertEqual(self.client.get(reverse("core:notification_open", args=[other.pk])).status_code, 404)
        self.client.post(reverse("core:notifications_read_all"))

    def test_notification_open_redirect_is_safe(self):
        n = Notification.objects.create(user=self.student, title="T", link="https://evil.example.com/")
        self.login(self.student)
        r = self.client.get(reverse("core:notification_open", args=[n.pk]))
        self.assertEqual(r["Location"], reverse("core:notifications"))

    def test_directions_json(self):
        data = self.client.get(reverse("core:directions_json"), {"faculty": self.f1.pk}).json()
        self.assertEqual([d["name"] for d in data["results"]], ["Buxgalteriya"])
        self.assertEqual(self.client.get(reverse("core:directions_json")).json()["results"], [])

    def test_unpublished_news_hidden(self):
        from core.models import News
        n = News.objects.create(title="Yashirin", body="b", is_published=False)
        self.assertEqual(self.client.get(reverse("core:news_detail", args=[n.pk])).status_code, 404)
        self.login(self.staff)
        self.assertEqual(self.client.get(reverse("core:news_detail", args=[n.pk])).status_code, 200)

    def test_protected_files(self):
        from django.core.files.base import ContentFile
        cert = self.p_student.certificates.create(title="IELTS")
        cert.file.save("c.pdf", ContentFile(b"%PDF-1.4 test"))
        url = reverse("students:file", args=["certificate", cert.pk])
        self.assertEqual(self.client.get(url).status_code, 302)  # kirmagan
        self.login(self.grad)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.login(self.student)
        self.assertEqual(self.client.get(url).status_code, 200)
        self.login(self.staff)
        self.assertEqual(self.client.get(url).status_code, 200)
        cert.file.delete()

    def test_logout_requires_post(self):
        self.login(self.student)
        self.assertEqual(self.client.get(reverse("accounts:logout")).status_code, 405)

    def test_csrf_enforced(self):
        from django.test import Client
        c = Client(enforce_csrf_checks=True)
        c.force_login(self.student)
        self.assertEqual(c.post(reverse("vacancies:apply", args=[self.vacancy.pk])).status_code, 403)

    def test_backup_command(self):
        import os
        import sqlite3
        import tempfile
        from unittest import mock
        from django.conf import settings as dj_settings
        from django.core.management import call_command
        with tempfile.TemporaryDirectory() as tmp:
            src_path = os.path.join(tmp, "src.sqlite3")
            con = sqlite3.connect(src_path)
            con.execute("CREATE TABLE t (x INTEGER)")
            con.execute("INSERT INTO t VALUES (42)")
            con.commit()
            con.close()
            backups = os.path.join(tmp, "backups")
            with mock.patch.dict(dj_settings.DATABASES["default"], {"NAME": src_path}), override_settings(BACKUP_DIR=backups):
                call_command("backup_db", stdout=io.StringIO())
            files = [f for f in os.listdir(backups) if f.startswith("db_")]
            self.assertEqual(len(files), 1)
            copy = sqlite3.connect(os.path.join(backups, files[0]))
            self.assertEqual(copy.execute("SELECT x FROM t").fetchone()[0], 42)
            copy.close()

    def test_import_students_command(self):
        import tempfile
        from django.core.management import call_command
        csv_text = "external_id,last_name,first_name,patronymic,faculty,direction,course,graduation_year,phone,email\n" \
                   "H100,Aliyev,Vali,,Iqtisodiyot,Buxgalteriya,3,,+998901234500,h100@x.uz\n" \
                   "H101,Karimova,Zuhra,,Iqtisodiyot,Buxgalteriya,,2024,+998901234501,h101@x.uz\n"
        with tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8") as f:
            f.write(csv_text)
        call_command("import_students", f.name, stdout=io.StringIO())
        self.assertEqual(User.objects.get(username="H100").role, Role.STUDENT)
        self.assertEqual(User.objects.get(username="H101").role, Role.GRADUATE)
        call_command("import_students", f.name, stdout=io.StringIO())  # qayta import — dublikat yaratmaydi
        self.assertEqual(StudentProfile.objects.filter(external_id="H100").count(), 1)
